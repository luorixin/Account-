from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor, as_completed
from io import BytesIO
import os
import re

from openpyxl import load_workbook

from app.classifier import (
    AccountClassification,
    account_type_short_code,
    normalize_existing_account_type,
)


OUTPUT_HEADERS = [
    "New Account Type",
    "Account Type Short",
    "Review Status",
    "Classification Confidence",
    "Classification Reason",
    "Evidence URLs",
]

DEFAULT_CLASSIFICATION_MAX_WORKERS = 4
MAX_CLASSIFICATION_MAX_WORKERS = 8
CLASSIFICATION_TARGET_ACCOUNT_TYPES = {
    "jv",
    "joint venture",
    "joint venture(jv)",
    "joint venture (jv)",
    "public entity",
    "private equity",
    "private equity investee",
}
LOCAL_PRIORITY_ACCOUNT_TYPES = {
    "State-owned Enterprise(SOE)",
    "Multinational Corporation(MNC)",
    "Private Enterprise(POE)",
}


def process_workbook(file_bytes: bytes, classifier, progress_callback=None) -> bytes:
    workbook = load_workbook(BytesIO(file_bytes))
    sheet = workbook.worksheets[0]
    account_name_col, account_type_col = _find_required_columns(sheet)
    output_start_col = sheet.max_column + 1
    target_rows = _classification_target_rows(sheet, account_name_col, account_type_col)
    total_target_rows = len(target_rows)
    if progress_callback:
        progress_callback(
            {
                "processed": 0,
                "total": total_target_rows,
                "current_account": "",
                "result": None,
            }
        )

    for offset, header in enumerate(OUTPUT_HEADERS):
        sheet.cell(row=1, column=output_start_col + offset).value = header

    classification_results = _classify_target_rows(
        target_rows,
        classifier,
        _classification_max_workers(),
        progress_callback=_progress_callback(progress_callback, total_target_rows),
    )

    for row in range(2, sheet.max_row + 1):
        account_name = _cell_text(sheet.cell(row=row, column=account_name_col).value)
        account_type = _cell_text(sheet.cell(row=row, column=account_type_col).value)
        if _should_classify_row(account_name, account_type):
            result = classification_results[row]
            values = [
                result.account_type_text,
                account_type_short_code(result.account_type_text),
                result.review_status,
                result.confidence,
                result.reason,
                result.evidence_url_text,
            ]
        else:
            normalized_account_type = normalize_existing_account_type(account_type)
            values = [
                normalized_account_type,
                account_type_short_code(normalized_account_type),
                "",
                "",
                "",
                "",
            ]

        for offset, value in enumerate(values):
            sheet.cell(row=row, column=output_start_col + offset).value = value

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _classify_target_rows(
    target_rows: list[tuple[int, str]],
    classifier,
    max_workers: int,
    progress_callback=None,
) -> dict[int, AccountClassification]:
    rows_by_account: dict[str, list[int]] = {}
    for row, account_name in target_rows:
        rows_by_account.setdefault(account_name, []).append(row)

    results_by_row: dict[int, AccountClassification] = {}
    if not rows_by_account:
        return results_by_row

    worker_count = min(max_workers, len(rows_by_account))
    with ThreadPoolExecutor(max_workers=worker_count) as executor:
        futures = {
            executor.submit(_classify_account, classifier, account_name): account_name
            for account_name in rows_by_account
        }
        for future in as_completed(futures):
            account_name = futures[future]
            result = future.result()
            for row in rows_by_account[account_name]:
                results_by_row[row] = result
                if progress_callback:
                    progress_callback(account_name, result)
    return results_by_row


def _classify_account(classifier, account_name: str) -> AccountClassification:
    try:
        return classifier.classify(account_name)
    except Exception as exc:
        return AccountClassification(
            account_types=["Other"],
            confidence="Low",
            reason=f"Classification failed for this account: {exc}",
            evidence_urls=[],
            review_status="Needs Review",
        )


def _progress_callback(progress_callback, total_target_rows: int):
    if not progress_callback:
        return None
    processed_target_rows = 0

    def report(account_name: str, result: AccountClassification):
        nonlocal processed_target_rows
        processed_target_rows += 1
        progress_callback(
            {
                "processed": processed_target_rows,
                "total": total_target_rows,
                "current_account": account_name,
                "result": result,
            }
        )

    return report


def _classification_max_workers() -> int:
    raw_value = os.getenv("CLASSIFICATION_MAX_WORKERS", "")
    try:
        max_workers = int(raw_value)
    except ValueError:
        max_workers = DEFAULT_CLASSIFICATION_MAX_WORKERS
    return max(1, min(max_workers, MAX_CLASSIFICATION_MAX_WORKERS))


def _find_required_columns(sheet):
    headers = {
        _normalize_header(sheet.cell(row=1, column=col).value): col
        for col in range(1, sheet.max_column + 1)
    }
    account_name_col = headers.get("account name")
    account_type_col = headers.get("account type")
    if not account_name_col or not account_type_col:
        raise ValueError("Missing required columns: Account Name and Account Type.")
    return account_name_col, account_type_col


def _classification_target_rows(sheet, account_name_col, account_type_col) -> list[tuple[int, str]]:
    rows = []
    for row in range(2, sheet.max_row + 1):
        account_name = _cell_text(sheet.cell(row=row, column=account_name_col).value)
        account_type = _cell_text(sheet.cell(row=row, column=account_type_col).value)
        if _should_classify_row(account_name, account_type):
            rows.append((row, account_name))
    return rows


def _should_classify_row(account_name: str, account_type: str) -> bool:
    normalized_account_type = normalize_existing_account_type(account_type)
    if normalized_account_type in LOCAL_PRIORITY_ACCOUNT_TYPES:
        return False
    return bool(account_name.strip()) and (
        account_type.strip() == ""
        or any(
            token in CLASSIFICATION_TARGET_ACCOUNT_TYPES
            for token in _account_type_tokens(account_type)
        )
    )


def _account_type_tokens(account_type: str) -> list[str]:
    return [
        token.strip().casefold()
        for token in re.split(r"[,;，；]", account_type)
        if token.strip()
    ]


def _normalize_header(value) -> str:
    return _cell_text(value).casefold()


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()
