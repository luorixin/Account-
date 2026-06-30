from io import BytesIO
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app.classifier import AccountClassification
from app.excel_processor import _should_classify_row, process_workbook


class FakeClassifier:
    def __init__(self):
        self.seen = []

    def classify(self, account_name):
        self.seen.append(account_name)
        return AccountClassification(
            account_types=["State-owned Enterprise(SOE)"],
            confidence="High",
            reason=f"{account_name} has public evidence of state ownership.",
            evidence_urls=["https://example.com/source"],
        )


class PartiallyFailingClassifier:
    def __init__(self):
        self.seen = []

    def classify(self, account_name):
        self.seen.append(account_name)
        if account_name == "Public Co":
            raise TimeoutError("LLM request timed out")
        return AccountClassification(
            account_types=["Multinational Corporation(MNC)"],
            confidence="Medium",
            reason="Fallback classification.",
            evidence_urls=[],
        )


class SlowClassifier:
    def __init__(self, delay_seconds=0.2):
        self.delay_seconds = delay_seconds
        self.seen = []

    def classify(self, account_name):
        time.sleep(self.delay_seconds)
        self.seen.append(account_name)
        return AccountClassification(
            account_types=[f"Other"],
            confidence="Medium",
            reason=f"Classified {account_name}.",
            evidence_urls=[],
        )


class VariableDelayClassifier:
    def __init__(self):
        self.seen = []

    def classify(self, account_name):
        self.seen.append(account_name)
        if account_name == "Slow Co":
            time.sleep(0.25)
        return AccountClassification(
            account_types=["Private Enterprise(POE)"],
            confidence="High",
            reason=f"{account_name} result.",
            evidence_urls=[],
        )


def workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Accounts"
    sheet.append(["Account Name", "Account Type", "Region"])
    sheet.append(["Public Co", "Public Entity", "CN"])
    sheet.append(["Blank Type Co", None, "JP"])
    sheet.append(["Whitespace Type Co", "   ", "DE"])
    sheet.append(["Private Equity Co", "Private Equity", "US"])
    sheet.append(["Private Equity Investee Co", "Private Equity Investee", "US"])
    sheet.append(["JV Co", "JV", "US"])
    sheet.append(["", None, "FR"])
    sheet.append(["Private Co", "Private Enterprise(POE)", "US"])
    other = workbook.create_sheet("Do Not Touch")
    other.append(["Account Name", "Account Type"])
    other.append(["Other Public", "Public Entity"])

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class ExcelProcessorTests(unittest.TestCase):
    def setUp(self):
        import app.db
        app.db.DB_PATH = "test_cleanse_tool.db"
        app.db.init_db()
        with app.db.get_db_connection() as conn:
            conn.execute("DELETE FROM classification_cache")
            conn.execute("DELETE FROM job_results")
            conn.execute("DELETE FROM jobs")
            conn.commit()

    def tearDown(self):
        import os
        if os.path.exists("test_cleanse_tool.db"):
            try:
                os.remove("test_cleanse_tool.db")
            except OSError:
                pass

    def test_processes_first_sheet_and_appends_classification_columns(self):
        classifier = FakeClassifier()

        output = process_workbook(workbook_bytes(), classifier)

        result = load_workbook(BytesIO(output))
        sheet = result["Accounts"]
        self.assertEqual(
            [sheet.cell(1, col).value for col in range(1, 10)],
            [
                "Account Name",
                "Account Type",
                "Region",
                "New Account Type",
                "Account Type Short",
                "Review Status",
                "Classification Confidence",
                "Classification Reason",
                "Evidence URLs",
            ],
        )
        self.assertEqual(sheet.cell(2, 4).value, "State-owned Enterprise(SOE)")
        self.assertEqual(sheet.cell(2, 5).value, "SOE")
        self.assertEqual(sheet.cell(2, 6).value, None)
        self.assertEqual(sheet.cell(2, 7).value, "High")
        self.assertIn("state ownership", sheet.cell(2, 8).value)
        self.assertEqual(sheet.cell(2, 9).value, "https://example.com/source")
        self.assertEqual(sheet.cell(3, 4).value, "State-owned Enterprise(SOE)")
        self.assertEqual(sheet.cell(3, 5).value, "SOE")
        self.assertEqual(sheet.cell(4, 4).value, "State-owned Enterprise(SOE)")
        self.assertEqual(sheet.cell(4, 5).value, "SOE")
        self.assertEqual(sheet.cell(5, 4).value, "State-owned Enterprise(SOE)")
        self.assertEqual(sheet.cell(5, 5).value, "SOE")
        self.assertEqual(sheet.cell(6, 4).value, "State-owned Enterprise(SOE)")
        self.assertEqual(sheet.cell(6, 5).value, "SOE")
        self.assertEqual(sheet.cell(7, 4).value, "State-owned Enterprise(SOE)")
        self.assertEqual(sheet.cell(7, 5).value, "SOE")
        self.assertEqual(sheet.cell(8, 4).value, None)
        self.assertEqual(sheet.cell(8, 5).value, None)
        self.assertEqual(sheet.cell(9, 4).value, "Private Enterprise(POE)")
        self.assertEqual(sheet.cell(9, 5).value, "POE")
        self.assertEqual(
            sorted(classifier.seen),
            [
                "Blank Type Co",
                "JV Co",
                "Private Equity Co",
                "Private Equity Investee Co",
                "Public Co",
                "Whitespace Type Co",
            ],
        )
        self.assertEqual(result["Do Not Touch"].max_column, 2)

    def test_missing_required_columns_raises_value_error(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Name", "Type"])
        stream = BytesIO()
        workbook.save(stream)

        with self.assertRaisesRegex(ValueError, "Account Name.*Account Type"):
            process_workbook(stream.getvalue(), FakeClassifier())

    def test_classifier_failure_marks_row_needs_review_and_continues(self):
        classifier = PartiallyFailingClassifier()

        output = process_workbook(workbook_bytes(), classifier)

        result = load_workbook(BytesIO(output))
        sheet = result["Accounts"]
        self.assertEqual(sheet.cell(2, 4).value, "Other")
        self.assertEqual(sheet.cell(2, 5).value, "Other")
        self.assertEqual(sheet.cell(2, 6).value, "Needs Review")
        self.assertEqual(sheet.cell(2, 7).value, "Low")
        self.assertIn("Classification failed for this account", sheet.cell(2, 8).value)
        self.assertEqual(sheet.cell(3, 4).value, "Multinational Corporation(MNC)")
        self.assertEqual(sheet.cell(3, 5).value, "MNC")
        self.assertEqual(sheet.cell(3, 6).value, None)
        self.assertEqual(sheet.cell(4, 4).value, "Multinational Corporation(MNC)")
        self.assertEqual(sheet.cell(4, 5).value, "MNC")
        self.assertEqual(sheet.cell(4, 6).value, None)
        self.assertEqual(sheet.cell(5, 4).value, "Multinational Corporation(MNC)")
        self.assertEqual(sheet.cell(5, 5).value, "MNC")
        self.assertEqual(sheet.cell(6, 4).value, "Multinational Corporation(MNC)")
        self.assertEqual(sheet.cell(6, 5).value, "MNC")
        self.assertEqual(sheet.cell(7, 4).value, "Multinational Corporation(MNC)")
        self.assertEqual(sheet.cell(7, 5).value, "MNC")
        self.assertEqual(sheet.cell(9, 4).value, "Private Enterprise(POE)")
        self.assertEqual(sheet.cell(9, 5).value, "POE")

    def test_progress_callback_reports_classification_target_progress(self):
        events = []

        process_workbook(workbook_bytes(), FakeClassifier(), progress_callback=events.append)

        self.assertEqual(len(events), 7)
        self.assertEqual(events[0]["processed"], 0)
        self.assertEqual(events[0]["total"], 6)
        self.assertEqual(events[1]["processed"], 1)
        self.assertEqual(events[1]["total"], 6)
        self.assertEqual(events[2]["processed"], 2)
        self.assertEqual(events[2]["total"], 6)
        self.assertEqual(events[3]["processed"], 3)
        self.assertEqual(events[3]["total"], 6)
        self.assertEqual(events[4]["processed"], 4)
        self.assertEqual(events[4]["total"], 6)
        self.assertEqual(events[5]["processed"], 5)
        self.assertEqual(events[5]["total"], 6)
        self.assertEqual(events[6]["processed"], 6)
        self.assertEqual(events[6]["total"], 6)
        self.assertEqual(
            {event["current_account"] for event in events[1:]},
            {
                "Public Co",
                "Blank Type Co",
                "Whitespace Type Co",
                "Private Equity Co",
                "Private Equity Investee Co",
                "JV Co",
            },
        )

    def test_target_account_type_matching_includes_private_equity_labels(self):
        cases = [
            ("Any Account", "", True),
            ("Any Account", "public entity", True),
            ("Any Account", "PRIVATE EQUITY", True),
            ("Any Account", "private equity investee", True),
            ("Any Account", "JV", True),
            ("Any Account", "Joint Venture", True),
            ("Any Account", "Joint Venture(JV)", True),
            ("Any Account", "Joint Venture (JV)", True),
            ("Any Account", "Private Equity,Public Entity", True),
            ("Any Account", "Private Equity, Public Entity", True),
            ("Any Account", "Joint Venture,Public Entity", True),
            ("Any Account", "Public Entity; Private Enterprise(POE)", False),
            ("Any Account", "Joint Venture, State-Owned Enterprise (SOE)", False),
            ("Any Account", "Public Entity, Multinational Corporation（MNC）", False),
            ("Any Account", "Private Equity, Private Entity（POE）", False),
            ("Any Account", "Private Enterprise(POE), Bank", False),
            ("Any Account", "Private-Equity", False),
            ("Any Account", "Private Equity Fund", False),
            ("", "JV", False),
            ("", "Private Equity", False),
        ]

        for account_name, account_type, expected in cases:
            with self.subTest(account_name=account_name, account_type=account_type):
                self.assertEqual(
                    _should_classify_row(account_name, account_type),
                    expected,
                )

    def test_normalizes_non_classification_account_types_locally(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Account Name", "Account Type"])
        sheet.append(["SE Co", "Specialized Enterprise (SE/央企)"])
        sheet.append(["SOE Co", "State-Owned Enterprise (SOE)"])
        sheet.append(["GO Co", "Governmental Organization (GO)"])
        sheet.append(["POE Co", "Private Entity（POE）"])
        sheet.append(["MNC Co", "Multinational Corporation（MNC）"])
        sheet.append(["NPO Co", "Non-Profit Organization (NPO)"])
        sheet.append(["Unknown Co", "Bank"])
        stream = BytesIO()
        workbook.save(stream)
        classifier = FakeClassifier()

        output = process_workbook(stream.getvalue(), classifier)

        result = load_workbook(BytesIO(output))
        output_values = [
            result.active.cell(row=row, column=3).value
            for row in range(2, result.active.max_row + 1)
        ]
        short_values = [
            result.active.cell(row=row, column=4).value
            for row in range(2, result.active.max_row + 1)
        ]
        self.assertEqual(
            output_values,
            [
                "State-owned Enterprise(SOE)",
                "State-owned Enterprise(SOE)",
                "State-owned Enterprise(SOE)",
                "Private Enterprise(POE)",
                "Multinational Corporation(MNC)",
                "Other",
                None,
            ],
        )
        self.assertEqual(
            short_values,
            [
                "SOE",
                "SOE",
                "SOE",
                "POE",
                "MNC",
                "Other",
                None,
            ],
        )
        self.assertEqual(classifier.seen, [])

    def test_multi_value_account_type_prefers_existing_priority_type_locally(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Account Name", "Account Type"])
        sheet.append(["SOE Mixed Co", "Joint Venture, State-Owned Enterprise (SOE)"])
        sheet.append(["MNC Mixed Co", "Public Entity, Multinational Corporation（MNC）"])
        sheet.append(["POE Mixed Co", "Private Equity, Private Entity（POE）"])
        stream = BytesIO()
        workbook.save(stream)
        classifier = FakeClassifier()

        output = process_workbook(stream.getvalue(), classifier)

        result = load_workbook(BytesIO(output))
        self.assertEqual(result.active.cell(2, 3).value, "State-owned Enterprise(SOE)")
        self.assertEqual(result.active.cell(2, 4).value, "SOE")
        self.assertEqual(result.active.cell(3, 3).value, "Multinational Corporation(MNC)")
        self.assertEqual(result.active.cell(3, 4).value, "MNC")
        self.assertEqual(result.active.cell(4, 3).value, "Private Enterprise(POE)")
        self.assertEqual(result.active.cell(4, 4).value, "POE")
        self.assertEqual(classifier.seen, [])

    def test_local_account_type_normalization_uses_priority_for_multiple_types(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Account Name", "Account Type"])
        sheet.append(["MNC Private Co", "Private Entity（POE）; Multinational Corporation（MNC）"])
        sheet.append(
            [
                "SOE MNC Co",
                "Joint Venture; State-Owned Enterprise (SOE); Multinational Corporation（MNC）",
            ]
        )
        stream = BytesIO()
        workbook.save(stream)

        output = process_workbook(stream.getvalue(), FakeClassifier())

        result = load_workbook(BytesIO(output))
        self.assertEqual(result.active.cell(2, 3).value, "Multinational Corporation(MNC)")
        self.assertEqual(result.active.cell(2, 4).value, "MNC")
        self.assertEqual(result.active.cell(3, 3).value, "State-owned Enterprise(SOE)")
        self.assertEqual(result.active.cell(3, 4).value, "SOE")

    def test_classifies_target_rows_concurrently(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Account Name", "Account Type"])
        for index in range(4):
            sheet.append([f"Public Co {index}", "Public Entity"])
        stream = BytesIO()
        workbook.save(stream)
        classifier = SlowClassifier(delay_seconds=0.2)

        started_at = time.perf_counter()
        process_workbook(stream.getvalue(), classifier)
        elapsed = time.perf_counter() - started_at

        self.assertLess(elapsed, 0.6)
        self.assertEqual(
            sorted(classifier.seen),
            ["Public Co 0", "Public Co 1", "Public Co 2", "Public Co 3"],
        )

    def test_writes_concurrent_results_back_to_original_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Account Name", "Account Type"])
        sheet.append(["Slow Co", "Public Entity"])
        sheet.append(["Fast Co", "Public Entity"])
        stream = BytesIO()
        workbook.save(stream)

        output = process_workbook(stream.getvalue(), VariableDelayClassifier())

        result = load_workbook(BytesIO(output))
        self.assertEqual(result.active.cell(2, 7).value, "Slow Co result.")
        self.assertEqual(result.active.cell(3, 7).value, "Fast Co result.")

    def test_reuses_classification_for_duplicate_account_names(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Account Name", "Account Type"])
        sheet.append(["Duplicate Co", "Public Entity"])
        sheet.append(["Duplicate Co", None])
        stream = BytesIO()
        workbook.save(stream)
        classifier = FakeClassifier()

        output = process_workbook(stream.getvalue(), classifier)

        result = load_workbook(BytesIO(output))
        self.assertEqual(classifier.seen, ["Duplicate Co"])
        self.assertEqual(result.active.cell(2, 3).value, "State-owned Enterprise(SOE)")
        self.assertEqual(result.active.cell(3, 3).value, "State-owned Enterprise(SOE)")

    def test_reuses_existing_output_columns_when_reprocessing_workbook(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Account Name", "Account Type", *[
            "New Account Type",
            "Account Type Short",
            "Review Status",
            "Classification Confidence",
            "Classification Reason",
            "Evidence URLs",
        ]])
        sheet.append([
            "Public Co",
            "Public Entity",
            "Old Type",
            "Old Short",
            "Old Review",
            "Old Confidence",
            "Old Reason",
            "Old URL",
        ])
        stream = BytesIO()
        workbook.save(stream)

        output = process_workbook(stream.getvalue(), FakeClassifier())

        result = load_workbook(BytesIO(output))
        sheet = result.active
        self.assertEqual(sheet.max_column, 8)
        self.assertEqual(sheet.cell(1, 3).value, "New Account Type")
        self.assertEqual(sheet.cell(2, 3).value, "State-owned Enterprise(SOE)")
        self.assertEqual(sheet.cell(2, 4).value, "SOE")
        self.assertEqual(sheet.cell(2, 8).value, "https://example.com/source")

    def test_classification_max_workers_can_force_serial_processing(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Account Name", "Account Type"])
        for index in range(2):
            sheet.append([f"Public Co {index}", "Public Entity"])
        stream = BytesIO()
        workbook.save(stream)

        with patch.dict("os.environ", {"CLASSIFICATION_MAX_WORKERS": "1"}):
            started_at = time.perf_counter()
            process_workbook(stream.getvalue(), SlowClassifier(delay_seconds=0.2))
            elapsed = time.perf_counter() - started_at

        self.assertGreaterEqual(elapsed, 0.4)


if __name__ == "__main__":
    unittest.main()
