from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from app.classifier import AccountClassification
from app.jobs import JobManager


class FakeClassifierFactory:
    def __call__(self, llm_api_key=None, tavily_api_key=None):
        return FakeClassifier()


class ReviewClassifierFactory:
    def __call__(self, llm_api_key=None, tavily_api_key=None):
        return ReviewClassifier()


class FakeClassifier:
    def classify(self, account_name):
        return AccountClassification(
            account_types=["Multinational Corporation(MNC)"],
            confidence="Medium",
            reason="LLM fallback - no public evidence found. Known multinational company.",
            evidence_urls=[],
        )


class ReviewClassifier:
    def classify(self, account_name):
        return AccountClassification(
            account_types=["Other"],
            confidence="Low",
            reason="Insufficient evidence; default recommendation needs review.",
            evidence_urls=[],
            review_status="Needs Review",
        )


def sample_workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Account Name", "Account Type"])
    sheet.append(["Mitsui & Co., Ltd.", "Public Entity"])
    sheet.append(["Tyson Foods, Inc.", None])
    sheet.append(["Existing Private Co", "Private Enterprise(POE)"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class JobManagerTests(unittest.TestCase):
    def test_job_reports_progress_and_exposes_download_when_complete(self):
        manager = JobManager(FakeClassifierFactory())

        job_id = manager.create_job(
            sample_workbook_bytes(),
            filename="accounts.xlsx",
            llm_api_key="test",
            tavily_api_key="",
            run_inline=True,
        )

        status = manager.get_status(job_id)
        self.assertEqual(status["state"], "completed")
        self.assertEqual(status["total"], 2)
        self.assertEqual(status["processed"], 2)
        self.assertEqual(status["success_count"], 2)
        self.assertEqual(status["needs_review_count"], 0)

        output, filename = manager.get_download(job_id)
        self.assertEqual(filename, "accounts_classified.xlsx")
        workbook = load_workbook(BytesIO(output))
        self.assertEqual(workbook.active.cell(2, 3).value, "Multinational Corporation(MNC)")
        self.assertEqual(workbook.active.cell(3, 3).value, "Multinational Corporation(MNC)")
        self.assertEqual(workbook.active.cell(4, 3).value, "Private Enterprise(POE)")

    def test_review_status_counts_separately_from_recommended_type(self):
        manager = JobManager(ReviewClassifierFactory())

        job_id = manager.create_job(
            sample_workbook_bytes(),
            filename="accounts.xlsx",
            llm_api_key="test",
            tavily_api_key="",
            run_inline=True,
        )

        status = manager.get_status(job_id)
        self.assertEqual(status["state"], "completed")
        self.assertEqual(status["success_count"], 2)
        self.assertEqual(status["needs_review_count"], 2)
        self.assertEqual(status["failure_count"], 0)

        output, _filename = manager.get_download(job_id)
        workbook = load_workbook(BytesIO(output))
        self.assertEqual(workbook.active.cell(2, 3).value, "Other")
        self.assertEqual(workbook.active.cell(2, 5).value, "Needs Review")


if __name__ == "__main__":
    unittest.main()
