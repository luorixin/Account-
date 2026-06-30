from io import BytesIO
import json
import socket
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app.classifier import AccountClassification
from app.jobs import JobManager
from app.server import _build_classifier, create_server


class FakeClassifierFactory:
    def __call__(self, llm_api_key=None, tavily_api_key=None):
        return FakeClassifier()


class FakeClassifier:
    def classify(self, account_name):
        return AccountClassification(
            account_types=["Private Enterprise(POE)"],
            confidence="Medium",
            reason="LLM fallback - no public evidence found. Name-based classification.",
            evidence_urls=[],
        )


def workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Account Name", "Account Type"])
    sheet.append(["Alibaba Group Holding Limited", "Public Entity"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class ServerJobApiTests(unittest.TestCase):
    def setUp(self):
        import app.db
        app.db.DB_PATH = "test_cleanse_tool.db"
        app.db.init_db()
        with app.db.get_db_connection() as conn:
            conn.execute("DELETE FROM classification_cache")
            conn.execute("DELETE FROM job_results")
            conn.execute("DELETE FROM jobs")
            conn.commit()

    def tearDown(self):
        import os
        if os.path.exists("test_cleanse_tool.db"):
            try:
                os.remove("test_cleanse_tool.db")
            except OSError:
                pass

    def test_index_page_serves_upload_ui(self):
        manager = JobManager(FakeClassifierFactory())
        server = create_server("127.0.0.1", 0, job_manager=manager)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base_url = f"http://127.0.0.1:{server.server_address[1]}/"
        try:
            with urllib.request.urlopen(base_url, timeout=5) as response:
                body = response.read().decode("utf-8")

            self.assertEqual(response.status, 200)
            self.assertIn("Account Type 清洗工具", body)
            self.assertIn("dropZone", body)
            self.assertIn("X-LLM-API-Key", body)
        finally:
            server.shutdown()
            server.server_close()

    def test_static_page_escapes_server_supplied_values_before_html_rendering(self):
        html = Path("app/static/index.html").read_text(encoding="utf-8")

        self.assertIn("function escapeHtml", html)
        self.assertIn("function safeHttpUrl", html)
        self.assertNotIn("${item.filename}", html)
        self.assertNotIn("${item.error}", html)
        self.assertNotIn("${row.reason}", html)
        self.assertNotIn("${row.account_name}", html)
        self.assertNotIn('href="${url}"', html)
        self.assertNotIn('value="${row.reason || ""}"', html)

    def test_build_classifier_enables_public_search_fallback_without_tavily_key(self):
        with (
            patch("app.server.WebSearchClient") as search_client,
            patch("app.server.OpenAICompatibleLlmClient") as llm_client,
        ):
            _build_classifier(llm_api_key="deepseek-test")

        search_client.assert_called_once_with(
            tavily_api_key=None,
            allow_public_fallback=True,
        )
        llm_client.assert_called_once_with(api_key="deepseek-test")

    def test_build_classifier_keeps_public_search_fallback_with_tavily_key(self):
        with (
            patch("app.server.WebSearchClient") as search_client,
            patch("app.server.OpenAICompatibleLlmClient") as llm_client,
        ):
            _build_classifier(llm_api_key="deepseek-test", tavily_api_key="tavily-test")

        search_client.assert_called_once_with(
            tavily_api_key="tavily-test",
            allow_public_fallback=True,
        )
        llm_client.assert_called_once_with(api_key="deepseek-test")

    def test_process_status_and_download_job_flow(self):
        manager = JobManager(FakeClassifierFactory())
        server = create_server("127.0.0.1", 0, job_manager=manager)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base_url = f"http://127.0.0.1:{server.server_address[1]}"
        try:
            request = urllib.request.Request(
                f"{base_url}/process",
                data=workbook_bytes(),
                headers={"X-Filename": "accounts.xlsx", "X-LLM-API-Key": "test"},
                method="POST",
            )
            with urllib.request.urlopen(request, timeout=5) as response:
                payload = json.loads(response.read().decode("utf-8"))
            self.assertIn("job_id", payload)

            status = _wait_for_complete(base_url, payload["job_id"])
            self.assertEqual(status["state"], "completed")
            self.assertEqual(status["processed"], 1)
            self.assertEqual(status["total"], 1)

            with urllib.request.urlopen(
                f"{base_url}/download?{urllib.parse.urlencode({'job_id': payload['job_id']})}",
                timeout=5,
            ) as response:
                output = response.read()
            workbook = load_workbook(BytesIO(output))
            self.assertEqual(workbook.active.cell(2, 3).value, "Private Enterprise(POE)")
        finally:
            server.shutdown()
            server.server_close()

    def test_history_results_and_manual_update_flow(self):
        manager = JobManager(FakeClassifierFactory())
        server = create_server("127.0.0.1", 0, job_manager=manager)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base_url = f"http://127.0.0.1:{server.server_address[1]}"
        try:
            request = urllib.request.Request(
                f"{base_url}/process",
                data=workbook_bytes(),
                headers={"X-Filename": "accounts.xlsx", "X-LLM-API-Key": "test"},
                method="POST",
            )
            with urllib.request.urlopen(request, timeout=5) as response:
                payload = json.loads(response.read().decode("utf-8"))
            job_id = payload["job_id"]
            _wait_for_complete(base_url, job_id)

            with urllib.request.urlopen(f"{base_url}/history", timeout=5) as response:
                history = json.loads(response.read().decode("utf-8"))
            self.assertEqual(history[0]["id"], job_id)
            self.assertEqual(history[0]["state"], "completed")

            with urllib.request.urlopen(
                f"{base_url}/job/results?{urllib.parse.urlencode({'job_id': job_id})}",
                timeout=5,
            ) as response:
                rows = json.loads(response.read().decode("utf-8"))
            self.assertEqual(rows[0]["row_index"], 2)
            self.assertEqual(rows[0]["new_account_type"], "Private Enterprise(POE)")

            update_payload = json.dumps(
                {
                    "job_id": job_id,
                    "row_index": 2,
                    "new_account_type": "Other",
                    "review_status": "",
                    "reason": "Manual correction.",
                }
            ).encode("utf-8")
            update_request = urllib.request.Request(
                f"{base_url}/job/update_result",
                data=update_payload,
                headers={"Content-Type": "application/json"},
                method="POST",
            )
            with urllib.request.urlopen(update_request, timeout=5) as response:
                update_result = json.loads(response.read().decode("utf-8"))
            self.assertTrue(update_result["success"])

            with urllib.request.urlopen(
                f"{base_url}/download?{urllib.parse.urlencode({'job_id': job_id})}",
                timeout=5,
            ) as response:
                output = response.read()
            workbook = load_workbook(BytesIO(output))
            self.assertEqual(workbook.active.cell(2, 3).value, "Other")
            self.assertEqual(workbook.active.cell(2, 7).value, "Manual correction.")
        finally:
            server.shutdown()
            server.server_close()

    def test_update_result_rejects_invalid_account_type(self):
        manager = JobManager(FakeClassifierFactory())
        server = create_server("127.0.0.1", 0, job_manager=manager)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base_url = f"http://127.0.0.1:{server.server_address[1]}"
        try:
            update_payload = json.dumps(
                {
                    "job_id": "job-1",
                    "row_index": 2,
                    "new_account_type": "<script>alert(1)</script>",
                    "review_status": "",
                    "reason": "",
                }
            ).encode("utf-8")
            update_request = urllib.request.Request(
                f"{base_url}/job/update_result",
                data=update_payload,
                headers={"Content-Type": "application/json"},
                method="POST",
            )

            with self.assertRaises(urllib.error.HTTPError) as ctx:
                urllib.request.urlopen(update_request, timeout=5)

            self.assertEqual(ctx.exception.code, 400)
            self.assertIn("Invalid new_account_type", ctx.exception.read().decode("utf-8"))
        finally:
            server.shutdown()
            server.server_close()

    def test_process_rejects_invalid_content_length_as_json_error(self):
        manager = JobManager(FakeClassifierFactory())
        server = create_server("127.0.0.1", 0, job_manager=manager)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        host, port = server.server_address
        try:
            with socket.create_connection((host, port), timeout=5) as client:
                client.sendall(
                    b"POST /process HTTP/1.1\r\n"
                    b"Host: 127.0.0.1\r\n"
                    b"Connection: close\r\n"
                    b"Content-Length: not-a-number\r\n"
                    b"X-Filename: accounts.xlsx\r\n"
                    b"\r\n"
                )
                client.shutdown(socket.SHUT_WR)
                response_chunks = []
                while True:
                    chunk = client.recv(4096)
                    if not chunk:
                        break
                    response_chunks.append(chunk)
                response = b"".join(response_chunks).decode("utf-8", errors="replace")

            self.assertIn("400", response)
            self.assertIn("Invalid Content-Length", response)
        finally:
            server.shutdown()
            server.server_close()


def _wait_for_complete(base_url, job_id):
    for _ in range(20):
        with urllib.request.urlopen(
            f"{base_url}/status?{urllib.parse.urlencode({'job_id': job_id})}",
            timeout=5,
        ) as response:
            status = json.loads(response.read().decode("utf-8"))
        if status["state"] in {"completed", "completed_with_errors", "failed"}:
            return status
        time.sleep(0.05)
    raise AssertionError("job did not complete")


if __name__ == "__main__":
    unittest.main()
